''' Open an excel file table and save it as a json file

json structure based on column names
    - dot '.' indicates dictionary structure
    - name under braket "[...]" indicates semicolon ";" separated multi-value data
    - name starting with hash '#' is ignored
'''
from openpyxl import load_workbook

from vbrtools.dicjsontools import merge_dict, save_json_file, create_nested_dict
from vbrtools.misctools import get_args
from vbrtools.misctools import force_stdout_encoding

def _exceltojson():
    ''' Entry point of this module
    '''
    force_stdout_encoding()
    args_def = [(['-f', '--inputfile'     ], {'action':'store', 'required':True, 'help':'Input Filename (xlsx) of the table',   'metavar':'xxx.xlsx'}),
                (['-w', '--inputworksheet'], {'action':'store', 'required':True, 'help':'Input worksheet name',                 'metavar':'xxx'     }),
                (['-t', '--inputtable'    ], {'action':'store', 'required':True, 'help':'Input table name',                     'metavar':'xxx'     }),
                (['-o', '--outputfile'    ], {'action':'store', 'required':True, 'help':'Output Filename (json)',               'metavar':'xxx.json'})
               ]
    args = get_args(args_def)

    wb = load_workbook(filename = args.inputfile, data_only=True)
    in_table = wb[args.inputworksheet].tables[args.inputtable]
    in_range = wb[args.inputworksheet][in_table.ref]
    column_names = in_table.column_names

    output = []
    for row in in_range[1:]: #skip first row which is the header
        output_entry = {}
        for name, rowCell in zip(column_names, row):
            cell_value = rowCell.value
            if (cell_value is not None) and (cell_value != '') and (name[0] != '#'):
                keys = name.split('.')

                # Detect multi value content
                if keys[-1][0] == '[' and keys[-1][-1] == ']':
                    keys[-1] = keys[-1][1:-1]
                    if str(cell_value).isnumeric():
                        cell_value = [cell_value]
                    else:
                        cell_value = str(cell_value).split(';')

                output_entry_key = create_nested_dict(keys, cell_value)
                output_entry = merge_dict(output_entry, output_entry_key)

        output.append(output_entry)

    save_json_file(output, args.outputfile)

if __name__ == "__main__":
    _exceltojson()
